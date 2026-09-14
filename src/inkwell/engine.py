from __future__ import annotations

import hashlib
import io
import math
import random
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont


FONT_DIR = Path(__file__).resolve().parent / "fonts"
PAPER_MM = {"a4": (210.0, 297.0), "letter": (215.9, 279.4), "a5": (148.0, 210.0)}
FONT_FILES = {
    "casual": "Kalam-Regular.ttf",
    "neat": "PatrickHand-Regular.ttf",
    "light": "Kalam-Light.ttf",
    "bold": "Kalam-Bold.ttf",
}
MAX_TEXT_LENGTH = 100_000


@dataclass(frozen=True)
class RenderOptions:
    paper: str = "a4"
    paper_style: str = "ruled"
    font: str = "casual"
    ink: str = "#173b67"
    font_size: float = 20.0
    variation: float = 0.32
    seed: str = "inkwell"
    dpi: int = 200
    margin_mm: float = 18.0
    line_spacing: float = 1.48
    title: str | None = None

    def validated(self) -> "RenderOptions":
        if self.paper not in PAPER_MM:
            raise ValueError(f"paper must be one of: {', '.join(PAPER_MM)}")
        if self.paper_style not in {"ruled", "grid", "blank"}:
            raise ValueError("paper_style must be ruled, grid, or blank")
        if self.font not in {"casual", "neat", "light"}:
            raise ValueError("font must be casual, neat, or light")
        if not re.fullmatch(r"#[0-9a-fA-F]{6}", self.ink):
            raise ValueError("ink must be a six-digit hex color")
        if not 10 <= self.font_size <= 42:
            raise ValueError("font_size must be from 10 to 42 points")
        if not 0 <= self.variation <= 1:
            raise ValueError("variation must be from 0 to 1")
        if not 96 <= self.dpi <= 300:
            raise ValueError("dpi must be from 96 to 300")
        if not 8 <= self.margin_mm <= 35:
            raise ValueError("margin_mm must be from 8 to 35")
        if not 1.15 <= self.line_spacing <= 2.0:
            raise ValueError("line_spacing must be from 1.15 to 2.0")
        if len(self.seed) > 200:
            raise ValueError("seed must be no longer than 200 characters")
        return self


@dataclass(frozen=True)
class TextLine:
    text: str
    kind: str = "body"
    indent: int = 0


@dataclass(frozen=True)
class Placement:
    page: int
    line: int
    text: str
    box: tuple[float, float, float, float]


@dataclass
class RenderedDocument:
    pages: list[Image.Image]
    placements: list[Placement]
    options: RenderOptions


def _px(mm: float, dpi: int) -> int:
    return round(mm * dpi / 25.4)


def _font_path(style: str, bold: bool = False) -> Path:
    name = "Kalam-Bold.ttf" if bold else FONT_FILES[style]
    path = FONT_DIR / name
    if not path.is_file():
        raise FileNotFoundError(f"Bundled font is missing: {path}")
    return path


def _rng(seed: str, scope: str) -> random.Random:
    digest = hashlib.sha256(f"{seed}\0{scope}".encode("utf-8")).digest()
    return random.Random(int.from_bytes(digest[:8], "big"))


def _parse_lines(text: str, title: str | None) -> list[TextLine]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    raw = normalized.split("\n")
    lines: list[TextLine] = []
    if title:
        lines.extend((TextLine(title.strip(), "title"), TextLine("", "blank")))
    for value in raw:
        clean = value.rstrip()
        if not clean.strip():
            lines.append(TextLine("", "blank"))
        elif clean.startswith("### "):
            lines.append(TextLine(clean[4:].strip(), "heading3"))
        elif clean.startswith("## "):
            lines.append(TextLine(clean[3:].strip(), "heading2"))
        elif clean.startswith("# "):
            lines.append(TextLine(clean[2:].strip(), "heading1"))
        elif re.match(r"^\s*[-*+]\s+", clean):
            content = re.sub(r"^\s*[-*+]\s+", "", clean)
            lines.append(TextLine(f"• {content}", "body", 1))
        elif re.match(r"^\s*\[[ xX]\]\s+", clean):
            mark = "☒" if re.match(r"^\s*\[[xX]\]", clean) else "☐"
            content = re.sub(r"^\s*\[[ xX]\]\s+", "", clean)
            lines.append(TextLine(f"{mark} {content}", "body", 1))
        else:
            lines.append(TextLine(clean.strip(), "body"))
    while lines and lines[-1].kind == "blank":
        lines.pop()
    return lines or [TextLine("", "blank")]


def _split_long_word(word: str, font: ImageFont.FreeTypeFont, width: float) -> list[str]:
    if font.getlength(word) <= width:
        return [word]
    chunks: list[str] = []
    current = ""
    for char in word:
        candidate = current + char
        if current and font.getlength(candidate) > width:
            chunks.append(current)
            current = char
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def _wrap(text: str, font: ImageFont.FreeTypeFont, width: float, safety: float = 0) -> list[str]:
    if not text:
        return [""]
    words: list[str] = []
    for word in text.split():
        words.extend(_split_long_word(word, font, width))
    result: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        candidate_width = font.getlength(candidate) + safety * len(candidate.split())
        if current and candidate_width > width:
            result.append(current)
            current = word
        else:
            current = candidate
    if current:
        result.append(current)
    return result


def _paper(size: tuple[int, int], options: RenderOptions, line_height: int, top: int) -> Image.Image:
    page = Image.new("RGB", size, "#fffdf7")
    draw = ImageDraw.Draw(page)
    width, height = size
    if options.paper_style == "grid":
        step = max(20, line_height)
        color = (211, 225, 229)
        for x in range(top, width, step):
            draw.line((x, 0, x, height), fill=color, width=1)
        for y in range(top, height, step):
            draw.line((0, y, width, y), fill=color, width=1)
    elif options.paper_style == "ruled":
        color = (195, 216, 225)
        for y in range(top + line_height, height, line_height):
            draw.line((0, y, width, y), fill=color, width=max(1, options.dpi // 180))
        margin_x = _px(14, options.dpi)
        draw.line((margin_x, 0, margin_x, height), fill=(235, 180, 180), width=max(1, options.dpi // 120))
    return page


def _word_layer(word: str, font: ImageFont.FreeTypeFont, ink: str, opacity: int) -> Image.Image:
    bbox = font.getbbox(word, stroke_width=0)
    width = max(1, math.ceil(bbox[2] - bbox[0]))
    height = max(1, math.ceil(bbox[3] - bbox[1]))
    pad = 10
    layer = Image.new("RGBA", (width + 2 * pad, height + 2 * pad), (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    rgb = tuple(int(ink[i : i + 2], 16) for i in (1, 3, 5))
    draw.text((pad - bbox[0], pad - bbox[1]), word, font=font, fill=(*rgb, opacity))
    return layer


def render_pages(text: str, options: RenderOptions | None = None) -> RenderedDocument:
    if not isinstance(text, str) or not text.strip():
        raise ValueError("text must be a non-empty string")
    if len(text) > MAX_TEXT_LENGTH:
        raise ValueError(f"text exceeds the {MAX_TEXT_LENGTH:,} character limit")
    options = (options or RenderOptions()).validated()
    width_mm, height_mm = PAPER_MM[options.paper]
    size = (_px(width_mm, options.dpi), _px(height_mm, options.dpi))
    margin = _px(options.margin_mm, options.dpi)
    font_px = round(options.font_size * options.dpi / 72)
    body_font = ImageFont.truetype(str(_font_path(options.font)), font_px)
    line_height = max(round(font_px * options.line_spacing), round(body_font.getbbox("Ag")[3] * 1.35))
    content_width = size[0] - 2 * margin
    top = margin
    bottom = size[1] - margin
    pages = [_paper(size, options, line_height, top)]
    placements: list[Placement] = []
    y = top
    logical_line = 0

    for source in _parse_lines(text, options.title):
        if source.kind == "blank":
            y += round(line_height * 0.62)
            continue
        scale = {"title": 1.6, "heading1": 1.38, "heading2": 1.22, "heading3": 1.1}.get(source.kind, 1.0)
        bold = source.kind in {"title", "heading1", "heading2"}
        current_font = ImageFont.truetype(str(_font_path(options.font, bold=bold)), round(font_px * scale))
        current_height = round(line_height * max(1.0, scale * 1.04))
        indent = round(line_height * 0.55 * source.indent)
        word_safety = 3 + 2 * options.variation * options.dpi / 96
        for wrapped in _wrap(source.text, current_font, content_width - indent, word_safety):
            if y + current_height > bottom:
                pages.append(_paper(size, options, line_height, top))
                y = top
            page_index = len(pages) - 1
            draw_x = margin + indent
            space_width = current_font.getlength(" ")
            words = wrapped.split(" ") if wrapped else []
            line_rng = _rng(options.seed, f"{logical_line}:{wrapped}")
            for word_index, word in enumerate(words):
                measured = current_font.getlength(word)
                max_rotation = 0.55 * options.variation
                angle = line_rng.uniform(-max_rotation, max_rotation)
                jitter_x = line_rng.uniform(-0.35, 0.35) * options.variation * options.dpi / 96
                jitter_y = line_rng.uniform(-1.15, 1.15) * options.variation * options.dpi / 96
                opacity = round(244 - line_rng.uniform(0, 16) * options.variation)
                layer = _word_layer(word, current_font, options.ink, opacity)
                if angle:
                    layer = layer.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True)
                bbox = layer.getbbox()
                visual_width = math.ceil(measured)
                if bbox:
                    cropped = layer.crop(bbox)
                    visual_width = cropped.width
                    paste_x = round(draw_x + jitter_x)
                    paste_y = round(y + jitter_y)
                    pages[page_index].paste(cropped, (paste_x, paste_y), cropped)
                    placements.append(Placement(page_index, logical_line, word, (paste_x, paste_y, paste_x + cropped.width, paste_y + cropped.height)))
                # Advance by the font's measured width plus a rotation/jitter safety gutter.
                draw_x += max(measured, visual_width) + space_width + word_safety
            if source.kind.startswith("heading") or source.kind == "title":
                underline_y = y + round(current_font.getbbox("Ag")[3] * 1.08)
                underline_rng = _rng(options.seed, f"underline:{logical_line}")
                ImageDraw.Draw(pages[page_index]).line(
                    (margin + indent, underline_y, min(size[0] - margin, draw_x), underline_y + underline_rng.choice((-1, 0, 1))),
                    fill=options.ink,
                    width=max(1, options.dpi // 140),
                )
            y += current_height
            logical_line += 1
        if source.kind.startswith("heading") or source.kind == "title":
            y += round(line_height * 0.22)
    return RenderedDocument(pages, placements, options)


def _save_pdf(document: RenderedDocument, output: Path) -> list[Path]:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    width_mm, height_mm = PAPER_MM[document.options.paper]
    width_pt, height_pt = width_mm * 72 / 25.4, height_mm * 72 / 25.4
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(output), pagesize=(width_pt, height_pt), pageCompression=1)
    for page in document.pages:
        stream = io.BytesIO()
        page.save(stream, "PNG", optimize=True)
        stream.seek(0)
        canvas.drawImage(ImageReader(stream), 0, 0, width_pt, height_pt, preserveAspectRatio=False)
        canvas.showPage()
    canvas.save()
    return [output]


def _save_docx(document: RenderedDocument, output: Path) -> list[Path]:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
    from docx.shared import Mm

    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    width_mm, height_mm = PAPER_MM[document.options.paper]
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(width_mm), Mm(height_mm)
    section.top_margin = section.bottom_margin = Mm(0)
    section.left_margin = section.right_margin = Mm(0)
    normal = doc.styles["Normal"]
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = 0
    with tempfile.TemporaryDirectory(prefix="inkwell-docx-") as temp:
        temp_path = Path(temp)
        for index, page in enumerate(document.pages):
            image_path = temp_path / f"page-{index + 1:02d}.png"
            page.save(image_path, "PNG", optimize=True, dpi=(document.options.dpi, document.options.dpi))
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.line_spacing = 1
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))
            if index < len(document.pages) - 1:
                run.add_break(WD_BREAK.PAGE)
        doc.save(output)
    return [output]


def generate_document(
    text: str,
    output: str | Path,
    fmt: str | None = None,
    options: RenderOptions | None = None,
) -> tuple[RenderedDocument, list[Path]]:
    destination = Path(output).expanduser().resolve()
    chosen = (fmt or destination.suffix.lstrip(".") or "pdf").lower()
    if chosen == "jpg" or chosen == "jpeg":
        raise ValueError("JPEG is intentionally unsupported because it degrades handwriting; use PNG")
    if chosen not in {"pdf", "docx", "png"}:
        raise ValueError("format must be pdf, docx, or png")
    document = render_pages(text, options)
    if chosen == "pdf":
        if destination.suffix.lower() != ".pdf":
            destination = destination.with_suffix(".pdf")
        files = _save_pdf(document, destination)
    elif chosen == "docx":
        if destination.suffix.lower() != ".docx":
            destination = destination.with_suffix(".docx")
        files = _save_docx(document, destination)
    else:
        directory = destination if not destination.suffix else destination.parent / destination.stem
        directory.mkdir(parents=True, exist_ok=True)
        files = []
        for index, page in enumerate(document.pages):
            path = directory / f"page-{index + 1:02d}.png"
            page.save(path, "PNG", optimize=True, dpi=(document.options.dpi, document.options.dpi))
            files.append(path)
    return document, files
